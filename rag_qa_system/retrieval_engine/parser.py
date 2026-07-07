# parser.py
import os
import pdfplumber
from pdf2image import convert_from_path
import pytesseract
from docx import Document as DocxDocument
from pptx import Presentation
# 注意：这里使用新版 langchain 的 Document 导入路径
from langchain_core.documents import Document

# ---------- 以下是解析各种文件的辅助函数 ----------

def _parse_txt(file_path: str) -> list[Document]:
    """解析 .txt 文件，整个文件当作一个 Document"""
    with open(file_path, 'r', encoding='utf-8') as f:
        text = f.read()
    return [Document(page_content=text, metadata={"source": file_path, "file_type": "txt"})]


def _parse_pdf(file_path: str) -> list[Document]:
    """解析 .pdf 文件，如果文字太少则自动启用 OCR"""
    docs = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text:
                docs.append(Document(
                    page_content=text,
                    metadata={"source": file_path, "page_number": i + 1, "file_type": "pdf"}
                ))

    # 判断是否为扫描件：提取的文字少于100个字符
    total_text = "".join([d.page_content for d in docs])
    if len(total_text.strip()) < 100:
        return _parse_pdf_ocr(file_path)
    return docs


def _parse_pdf_ocr(file_path: str) -> list[Document]:
    """使用 OCR 解析图片型 PDF"""
    # Windows 用户务必修改下面这一行，改为你的 tesseract.exe 安装路径
    # pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

    images = convert_from_path(file_path, dpi=200)
    docs = []
    for i, img in enumerate(images):
        text = pytesseract.image_to_string(img, lang='chi_sim+eng')
        docs.append(Document(
            page_content=text,
            metadata={"source": file_path, "page_number": i + 1, "file_type": "pdf_ocr"}
        ))
    return docs


def _parse_docx(file_path: str) -> list[Document]:
    """解析 .docx 文件，所有段落合并成一个 Document"""
    doc = DocxDocument(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return [Document(page_content="\n".join(full_text), metadata={"source": file_path, "file_type": "docx"})]


def _parse_pptx(file_path: str) -> list[Document]:
    """解析 .pptx 文件，每一页幻灯片作为一个 Document"""
    prs = Presentation(file_path)
    docs = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    slide_text.append(para.text)
        docs.append(Document(
            page_content="\n".join(slide_text),
            metadata={"source": file_path, "slide_number": i + 1, "file_type": "pptx"}
        ))
    return docs


# ---------- 统一入口函数 ----------

def parse_document(file_path: str) -> list[Document]:
    """根据文件后缀调用对应的解析器，返回 Document 列表"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.txt':
        return _parse_txt(file_path)
    elif ext == '.pdf':
        return _parse_pdf(file_path)
    elif ext == '.docx':
        return _parse_docx(file_path)
    elif ext == '.pptx':
        return _parse_pptx(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")